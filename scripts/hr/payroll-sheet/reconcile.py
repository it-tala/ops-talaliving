#!/usr/bin/env python3
"""Run a signed weekly payroll sheet through ops, on a throwaway database, and
say — person by person — whether ops pays the same.

    supabase/local/rebuild.sh
    python3 scripts/hr/payroll-sheet/reconcile.py "9 ALL DAILY WORKER PAYROLL - WEEK 1 SEPTEMBER.xlsx"

## What it proves

The owner's question (D340): *if ops is given the same attendance and the same
overtime hours as the sheet, does it compute the same pay?* So the inputs come
from the sheet and the machine, and only the arithmetic is ops':

  - **taps**: the sheet's own BIOMETRIC paste (the machine's export), through
    `ops_hr.import_scans`, exactly as HRD uploads it;
  - **which days count**: the sheet's day columns (SENIN..JUMAT, SABTU for the
    guards) and its SABTU/MINGGU/TANGGAL MERAH count. Where the taps say
    otherwise the day is marked (absent, half day) or typed in, each with a
    reason — the decisions HRD makes on /hrd/absensi;
  - **overtime hours**: the sheet's own totals, 1,5× (U) and 2× (X), on one
    approved production sheet per week;
  - **saldo and potongan**: the sheet's columns, as run adjustments.

Then `ops_hr.period_lines` computes the week and each person's net is
compared with the sheet's TOTAL GAJI, component by component.

## Where it runs

**Only on a local scratch cluster** — it writes a rule book, employees, taps,
marks, overtime sheets and a run, with two test identities (`hrd-uji`,
`pimpinan-uji`) that exist nowhere else. It refuses any other host, the same
way `rebuild.sh` does. Nothing here touches production, and nothing here is
evidence that anybody approved anything: it is a calculator check.

Inputs other than the workbook: `machine_map.csv` (payroll name → machine
number, taken from the machine's own export, not the payroll's FP column —
F197) and the rule book in `RULES` below.
"""
import csv, datetime as dt, json, math, os, subprocess, sys, tempfile

try:
    import openpyxl
except ImportError:
    sys.exit("pip install openpyxl")

HERE = os.path.dirname(os.path.abspath(__file__))
HOST = os.environ.get("PGHOST", "/tmp")
PORT = os.environ.get("PGPORT", "5433")
USER = os.environ.get("PGUSER", "postgres")
if not (HOST.startswith("/") or HOST in ("localhost", "127.0.0.1", "::1")):
    sys.exit(f"refusing: PGHOST={HOST} is not local. This writes test identities and a payroll run.")

HRD = "ffffffff-0000-0000-0000-00000000d340"
LEAD = "ffffffff-0000-0000-0000-00000000d341"

# What a guard's Sunday is worth: 2×, the owner's ruling (D341) for any
# Monday–Saturday pattern. The September sheets paid three guards' Sundays as
# 2 units at 2× (F198); `SATPAM_SUNDAY=4` reproduces those sheets exactly.
SATPAM_SUNDAY = float(os.environ.get("SATPAM_SUNDAY", "2"))

# The rule book being checked: production's v5 patterns with D340's keys.
RULES = {
    "late_mode": "manual", "late_grace_minutes": 15, "undertime_mode": "off", "undertime_grace_minutes": 15,
    "week_pattern": "5day", "day_starts_minutes": 480, "effective_days_per_year": 240,
    "hourly_basis": "company", "monthly_divisor": 173, "flat_multiplier": 1, "late_forfeits_allowance": False,
    "overtime_mode": "statutory", "overtime_rounding_minutes": 0,
    "workday_tiers": [{"after_hours": 0, "multiplier": 1.5}],
    "restday_tiers": [{"after_hours": 0, "multiplier": 2}, {"after_hours": 8, "multiplier": 3}, {"after_hours": 9, "multiplier": 4}],
    # D340
    "day_reading": "schedule", "hours_rounding_minutes": 15, "out_window_minutes": 30,
    "holiday_pay_multiplier": 2, "allowance_on_premium_days": False, "allowance_by_day_value": True,
    "overtime_night_after_minutes": 1320, "overtime_night_multiplier": 2,
    "overtime_exact_hourly": True, "hourly_includes_allowance": False, "pay_week_starts_isodow": 6,
    "schedules": [
        {"code": "PRODUKSI", "name": "Produksi", "start_minutes": 450, "end_minutes": 990, "break_minutes": 45,
         "friday_end_minutes": 960, "friday_break_minutes": 90, "note": None,
         "days": {"6": {"start_minutes": 480, "end_minutes": 960, "break_minutes": 0, "pay_multiplier": 2},
                  "7": {"start_minutes": 480, "end_minutes": 960, "break_minutes": 0, "pay_multiplier": 2}}},
        # The sheet counts SABTU as an ordinary day for the helper, as for the
        # guards (its day total adds column F for both).
        {"code": "HELPER", "name": "Helper", "start_minutes": 450, "end_minutes": 990, "break_minutes": 45,
         "friday_end_minutes": 960, "friday_break_minutes": 90, "note": None,
         "days": {"6": {"start_minutes": 480, "end_minutes": 960, "break_minutes": 0, "pay_multiplier": 1},
                  "7": {"start_minutes": 480, "end_minutes": 960, "break_minutes": 0, "pay_multiplier": 2}}},
        {"code": "SATPAM", "name": "Satpam", "start_minutes": 1140, "end_minutes": 420, "break_minutes": 0,
         "friday_end_minutes": None, "friday_break_minutes": None, "note": None,
         "days": {"7": {"pay_multiplier": SATPAM_SUNDAY}}},
    ],
    "schedule_by_unit": {"Workshop": "PRODUKSI"},
}

def col(c):
    n = 0
    for ch in c: n = n * 26 + ord(ch) - 64
    return n - 1

def num(v):
    try: return float(v)
    except (TypeError, ValueError): return 0.0

def q(s): return "'" + str(s).replace("'", "''") + "'"

# ── the workbook ─────────────────────────────────────────────────────────
def read_book(path):
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    pay = next(wb[n] for n in wb.sheetnames if n.upper().startswith("ADJUSTED OLD FORMAT"))
    bio = next(wb[n] for n in wb.sheetnames if n.upper().startswith("PASTE HERE BIOMETRIC"))
    rows = [list(r) for r in pay.iter_rows(values_only=True)]
    # The row whose F..L are the seven dates, SABTU..JUMAT.
    hdr = next(i for i, r in enumerate(rows) if len(r) > 11 and isinstance(r[col('F')], dt.datetime)
               and isinstance(r[col('L')], dt.datetime))
    days = [rows[hdr][col(c)].date() for c in "FGHIJKL"]
    people = []
    for r in rows[hdr + 1:]:
        if len(r) <= col('AG') or not r[col('D')] or not isinstance(r[col('B')], (int, float)): continue
        g = lambda c: r[col(c)]
        people.append(dict(
            no=int(g('B')), name=str(g('D')).strip(), division=str(g('E') or '').strip(), rate=num(g('M')),
            day={days[i]: num(g(c)) for i, c in enumerate("FGHIJKL")},
            incentive=num(g('Q')), allowance=num(g('AC')), ot15=num(g('U')), ot2=num(g('X')),
            weekend=num(g('AA')), deduction=num(g('AE')), saldo=num(g('AF')), total=num(g('AG')),
            basic=num(g('P')), weekend_amt=num(g('AB')), ot_amt=num(g('V')) + num(g('Y')),
            allow_amt=num(g('S')) + num(g('AD')),
        ))
    taps = []
    for r in bio.iter_rows(values_only=True):
        if len(r) > 6 and isinstance(r[2], (int, float)) and isinstance(r[3], dt.datetime):
            taps.append((str(int(r[2])), r[3], (r[6] or 'FP')))
    return days, people, taps

# ── the database ─────────────────────────────────────────────────────────
def psql(sql, capture=False):
    with tempfile.NamedTemporaryFile("w", suffix=".sql", delete=False) as f:
        f.write("\\set ON_ERROR_STOP on\n" + sql)
        name = f.name
    try:
        res = subprocess.run(["psql", "-h", HOST, "-p", PORT, "-U", USER, "-q", "-X", "-A", "-t", "-F", "\t", "-f", name],
                             capture_output=True, text=True)
    finally:
        os.unlink(name)
    if res.returncode != 0:
        sys.exit("psql failed:\n" + res.stderr[-3000:])
    return res.stdout

def main(path):
    days, people, taps = read_book(path)
    start, end = days[0], days[-1]
    mmap = {r['payroll_name']: r for r in csv.DictReader(open(os.path.join(HERE, "machine_map.csv")))}
    missing = [p['name'] for p in people if p['name'] not in mmap]
    if missing:
        sys.exit("not in machine_map.csv: " + ", ".join(missing))
    # A rate in half rupiah (the guards' 52.083,5) is rounded up: D342.
    for p in people:
        m = mmap[p['name']]
        p['emp'] = m['machine_no'] or ("NM-" + p['name'].replace(' ', '_'))
        p['guard'] = 'SECURITY' in p['division'].upper()
        p['helper'] = 'HELPER' in p['division'].upper()
        p['pattern'] = 'SATPAM' if p['guard'] else ('HELPER' if p['helper'] else 'PRODUKSI')
    known = {p['emp'] for p in people}
    stray = sorted({t[0] for t in taps} - known)

    # 1. identities, rule book, employees, the machine's taps
    emp_rows = ",\n".join(
        f"  ({q(p['emp'])}, {q(p['name'])}, {q(p['division'])}, {math.floor(p['rate'] + 0.5)}, "
        f"{round(p['incentive'] + p['allowance'])}, {12 if p['guard'] else 8}, {q(p['pattern'])})"
        for p in people)
    tap_json = json.dumps([{"employee_ref": n, "at": at.strftime("%Y-%m-%dT%H:%M:%S") + "+07:00", "verify": v}
                           for n, at, v in taps if n in known])
    psql(f"""
do $$ begin
  if exists (select 1 from ops_hr.payroll_runs) or exists (select 1 from ops_hr.employees) then
    raise exception 'this database is not empty — run supabase/local/rebuild.sh first';
  end if;
end $$;
insert into auth.users (id, email, raw_user_meta_data) values
  ({q(HRD)}, 'hrd-uji@local.test', '{{"full_name":"HRD (uji rekonsiliasi)"}}'),
  ({q(LEAD)}, 'pimpinan-uji@local.test', '{{"full_name":"Pimpinan (uji rekonsiliasi)"}}');
insert into ops_core.user_modules (user_id, module, level) values
  ({q(HRD)}, 'hrd', 'admin'), ({q(HRD)}, 'payroll', 'admin');
insert into ops_hr.pay_rule_sets (version, effective_from, note, rules, created_by)
values (1, {q(start)}, 'uji rekonsiliasi', {q(json.dumps(RULES))}::jsonb, {q(HRD)});
insert into ops_hr.employees (employee_no, full_name, position, unit, pay_basis, base_rate, allowance_rate,
                              daily_hours, joined_on, paid_leave_days, active, schedule_code)
select v.no, v.n, v.pos, 'Workshop', 'daily', v.rate, v.allow, v.h, date '2026-01-01', 0, true, v.sc
  from (values
{emp_rows}
  ) v(no, n, pos, rate, allow, h, sc);
set role authenticated;
set request.jwt.claim.sub = {q(HRD)};
do $$
declare r jsonb;
begin
  r := ops_hr.import_scans('uji rekonsiliasi.csv', {q(tap_json)}::jsonb);
  if r ->> 'outcome' <> 'ok' then raise exception 'import: %', r; end if;
end $$;
""")

    # 2. what ops reads from those taps
    out = psql(f"""
select e.employee_no, d.work_date, d.day_value, d.taps
  from ops_hr.employees e
  cross join lateral ops_hr.timesheet(e.id, {q(start)}, {q(end)}) d;
""")
    read = {}
    for line in out.strip().splitlines():
        no, d, v, n = line.split("\t")
        read[(no, dt.date.fromisoformat(d))] = (float(v), int(n))

    # 3. the sheet's decisions where the taps differ
    new_taps, marks, notes = [], [], []
    def clock(d, guard):
        dow = d.isoweekday()
        if guard: return [(d, "19:00"), (d + dt.timedelta(days=1), "07:00")]
        if dow >= 6: return [(d, "08:00"), (d, "16:00")]
        return [(d, "07:30"), (d, "16:00" if dow == 5 else "16:30")]
    def want(p, d, v):
        has, n = read[(p['emp'], d)]
        if v == 0 and has > 0:
            marks.append((p['emp'], d, 'absent', 'Tidak dibayar di payroll sheet'))
        elif v > 0:
            if n == 0:
                for day, hhmm in clock(d, p['guard']):
                    new_taps.append((p['emp'], f"{day} {hhmm}", 'Hadir menurut payroll sheet — tidak ada tap di mesin'))
            if v == 0.5:
                marks.append((p['emp'], d, 'half_day', 'Setengah hari menurut payroll sheet'))
    for p in people:
        sat_ordinary = p['guard'] or p['helper']
        weekdays = [d for d in days if d.isoweekday() <= 5] + ([d for d in days if d.isoweekday() == 6] if sat_ordinary else [])
        for d in weekdays:
            want(p, d, p['day'][d])
        # SABTU/MINGGU/TANGGAL MERAH: a count, not days. Filled in date order,
        # days with taps first; what is left over is typed in or cannot be.
        # A guard's Sunday at 4× is one day that the sheet counts as 2 units.
        per_day = SATPAM_SUNDAY / 2 if p['guard'] else 1
        left = p['weekend'] / per_day
        wk = [d for d in days if d.isoweekday() == 7] if sat_ordinary else [d for d in days if d.isoweekday() >= 6]
        wk.sort(key=lambda d: (read[(p['emp'], d)][1] == 0, d))
        for d in wk:
            v = 1.0 if left >= 1 else (0.5 if left >= 0.5 else 0.0)
            want(p, d, v); left -= v
        if left > 0:
            notes.append(f"{p['name']}: the sheet pays {p['weekend']:g} weekend/red days at 2×; "
                         f"the week has {len(wk)} — {left:g} cannot be a day in ops")

    tap_sql = ",\n".join(f"  ({q(n)}, {q(a)}, {q(w)})" for n, a, w in new_taps) or None
    mark_sql = ",\n".join(f"  ({q(n)}, {q(d)}::date, {q(k)}::ops_hr.day_mark_t, {q(w)})" for n, d, k, w in marks) or None
    friday = next(d for d in days if d.isoweekday() == 5)
    ot = [p for p in people if p['ot15'] or p['ot2']]
    # The 2× hours are the part after 22.00, so the line finishes that long
    # after 22.00 — past midnight it is the next morning, as ops reads it.
    ot_sql = ",\n".join(
        f"  ({q(p['emp'])}, {p['ot15'] + p['ot2']}, {('null' if not p['ot2'] else int(1320 + p['ot2'] * 60) % 1440)}::int)"
        for p in ot)
    adj = []
    for p in people:
        if p['saldo']: adj.append((p['emp'], 'carry_over', round(p['saldo'], 2), 'SALDO payroll sheet'))
        if p['deduction']: adj.append((p['emp'], 'other', -round(p['deduction'], 2), 'LATE/Deducted payroll sheet'))
    adj_sql = ",\n".join(f"    ({q(n)}, {q(k)}::ops_hr.adjustment_kind_t, {a}, {q(w)})" for n, k, a, w in adj)

    psql(f"""
reset role;
{"" if not tap_sql else f'''insert into ops_hr.attendance_scans (employee_id, work_date, at, verify, source, reason, recorded_by)
select e.id, ops_core.office_day(t.at), t.at, 'MANUAL', 'manual', v.why, {q(HRD)}::uuid
  from (values
{tap_sql}
  ) v(no, at, why)
  join ops_hr.employees e on e.employee_no = v.no
  cross join lateral (select (v.at::timestamp) at time zone ops_core.office_tz() as at) t;'''}
{"" if not mark_sql else f'''insert into ops_hr.day_marks (employee_id, work_date, kind, reason, marked_by)
select e.id, v.d, v.k, v.why, {q(HRD)}::uuid from (values
{mark_sql}
  ) v(no, d, k, why) join ops_hr.employees e on e.employee_no = v.no;'''}
with s as (
  insert into ops_hr.overtime_sheets (kind, work_date, purpose, created_by,
      hrd_checked_by, hrd_checked_at, leader_approved_by, leader_approved_at)
  values ('production', {q(friday)}, 'Uji rekonsiliasi: total jam lembur payroll sheet (U 1,5× dan X 2×)',
      {q(HRD)}, {q(HRD)}, now(), {q(LEAD)}, now())
  returning id)
insert into ops_hr.overtime_lines (sheet_id, employee_id, hours, task, until_minutes)
select s.id, e.id, v.h, 'jam lembur payroll sheet', v.u from s, (values
{ot_sql}
  ) v(no, h, u) join ops_hr.employees e on e.employee_no = v.no;
set role authenticated;
set request.jwt.claim.sub = {q(HRD)};
do $$
declare r jsonb; v_run text; x record;
begin
  r := ops_hr.open_payroll_run({q(start)}, {q(end)}, 'uji rekonsiliasi');
  if r ->> 'outcome' <> 'ok' then raise exception 'open run: %', r; end if;
  v_run := r -> 'data' ->> 'run_no';
  for x in select * from (values
{adj_sql}
  ) v(no, k, a, why) loop
    r := ops_hr.add_adjustment(v_run, x.no, x.k, x.a, x.why);
    if r ->> 'outcome' <> 'ok' then raise exception 'adjustment %: %', x.no, r; end if;
  end loop;
end $$;
""")

    # 4. what ops pays, against the sheet
    out = psql("""
select l.employee_no, l.base_pay, l.allowance_pay, l.overtime_pay, l.adjustment_total, l.net, l.worked_days, l.open_days
  from ops_hr.payroll_runs r
  cross join lateral ops_hr.period_lines(r.period_start, r.period_end, r.run_no) l;
""")
    ops = {}
    for line in out.strip().splitlines():
        f = line.split("\t")
        ops[f[0]] = dict(base=int(f[1]), allow=int(f[2]), ot=int(f[3]), adj=float(f[4]), net=float(f[5]),
                         days=float(f[6]), open=int(f[7]))

    rows = []
    tot_s = tot_o = 0.0
    for p in people:
        o = ops[p['emp']]
        s_base = p['basic'] + p['weekend_amt']
        s_adj = p['saldo'] - p['deduction']
        rows.append([p['no'], p['name'], p['emp'],
                     round(s_base, 2), o['base'], round(p['allow_amt'], 2), o['allow'],
                     round(p['ot_amt'], 2), o['ot'], round(s_adj, 2), o['adj'],
                     round(p['total'], 2), o['net'], round(o['net'] - p['total'], 2)])
        tot_s += p['total']; tot_o += o['net']
    hdr = ["no", "nama", "no_ops", "sheet_upah", "ops_upah", "sheet_tunjangan", "ops_tunjangan",
           "sheet_lembur", "ops_lembur", "sheet_saldo_potongan", "ops_penyesuaian", "sheet_total", "ops_net", "selisih"]
    dest = os.environ.get("OUT", "reconcile.csv")
    with open(dest, "w", newline="") as f:
        csv.writer(f).writerows([hdr] + rows)

    print(f"Periode {start} – {end} · {len(people)} orang · {len(taps)} tap mesin · "
          f"{len(new_taps)} tap diketik · {len(marks)} tanda hari")
    print(f"{'no':>3} {'nama':<12} {'sheet':>14} {'ops':>12} {'selisih':>10}  komponen yang beda")
    for r in rows:
        parts = [n for n, a, b in (("upah", r[3], r[4]), ("tunjangan", r[5], r[6]), ("lembur", r[7], r[8]),
                                   ("penyesuaian", r[9], r[10])) if abs(a - b) >= 1]
        flag = "" if abs(r[13]) < 1 else "  ← " + ", ".join(f"{n} {b - a:+,.2f}" for n, a, b in
                  (("upah", r[3], r[4]), ("tunjangan", r[5], r[6]), ("lembur", r[7], r[8]),
                   ("penyesuaian", r[9], r[10])) if abs(a - b) >= 1)
        print(f"{r[0]:>3} {r[1][:12]:<12} {r[11]:>14,.2f} {r[12]:>12,.0f} {r[13]:>10,.2f}{flag}")
    print(f"    {'TOTAL':<12} {tot_s:>14,.2f} {tot_o:>12,.0f} {tot_o - tot_s:>10,.2f}")
    # The sheet's attendance decisions ops needed on top of the machine — HRD's
    # work on /hrd/absensi, listed so it can be checked against the paper.
    by = {}
    for n, at, _ in new_taps: by.setdefault(n, set()).add(("diketik", at[:10]))
    for n, d, k, _ in marks: by.setdefault(n, set()).add((k, str(d)))
    names = {p['emp']: p['name'] for p in people}
    if by:
        print("keputusan hari dari sheet (bukan dari mesin):")
        for n in sorted(by, key=lambda n: names[n]):
            print(f"  {names[n]:<12} " + ", ".join(f"{d[5:]} {k}" for k, d in sorted(by[n], key=lambda x: x[1])))
    for n in notes: print("catatan:", n)
    if stray: print("nomor mesin di export yang tidak ada di payroll:", ", ".join(stray))
    print("per orang:", dest)

if __name__ == "__main__":
    if len(sys.argv) != 2: sys.exit(__doc__)
    main(sys.argv[1])
